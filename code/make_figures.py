#!/usr/bin/env python3
"""Build the data behind Figures 2-6 and S1-S3, write it to figure_data.xlsx, and draw the figures.

    python3 code/make_figures.py

If output/ exists (after `Rscript code/run_analysis.R`), figure data are built from that fresh run
and written to output/figure_data.xlsx. Otherwise the shipped results/figure_data.xlsx is used,
so the figures can be redrawn without R. Every plotted value comes from the figure-data sheets.
Figure 1 is a conceptual schematic with no underlying data.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
RUN = ROOT / "output"
SHIPPED = ROOT / "results" / "figure_data.xlsx"
TT = 25

SHEETS = {
    "Figure2": "Figure 2. Fuel carbon intensity (g CO2/MJ): global central case, reference (delay 0) and catch-up after 3-, 5- and 10-year delays",
    "Figure3a": "Figure 3a. Cumulative penalty (Gt CO2) by delay, global central case",
    "Figure3b": "Figure 3b. Low-carbon-fuel mass-equivalent supply (Mt/yr at 44 MJ/kg): reference (delay 0) and catch-up",
    "Figure4": "Figure 4. Annual emissions (Gt CO2/yr) and cumulative emissions (Gt CO2): reference (delay 0) and five-year delay",
    "Figure5": "Figure 5. Catch-up penalty (Gt CO2) under demand and energy-intensity (efficiency) variants",
    "Figure6_curves": "Figure 6a. Cumulative emissions overshoot after a five-year delay (Gt CO2; % of reference cumulative emissions to date)",
    "Figure6_gates": "Figure 6b. Review-gate diagnostics at 2030, 2035 and 2040 (full precision of Table S14; peaks in Mt fuel equivalent/yr)",
    "FigureS1": "Figure S1. Australian cumulative penalty (Mt CO2) by delay, BITRE-AES proxy (catch-up infeasible at 25 years)",
    "FigureS2": "Figure S2. Australian catch-up penalty (Mt CO2), harmonized BAU, IP and AR demand",
    "FigureS3a": "Figure S3a. Reference transition shapes: fuel carbon intensity (g CO2/MJ) and normalized remaining intensity",
    "FigureS3b": "Figure S3b. Catch-up penalty (Gt CO2) by transition shape",
}


# ---------------------------------------------------------------------------
# Figure data
# ---------------------------------------------------------------------------
def build_from_run() -> dict[str, pd.DataFrame]:
    ann = pd.read_csv(RUN / "annual_results.csv", float_precision="round_trip")
    diag = pd.read_csv(RUN / "diagnostics.csv", float_precision="round_trip")
    gates = pd.read_csv(RUN / "tables" / "TableS14_review_gates.csv", float_precision="round_trip")
    full_a = ann[ann["mode"] == "Full pause"]
    full_d = diag[diag["mode"] == "Full pause"]
    g = full_a[full_a.reference_id == "Global:Demand-IP"]
    cu = g[(g.response == "Catch-up") & g.delay.isin([0, 3, 5, 10])].sort_values(["delay", "year"])

    f = {}
    f["Figure2"] = cu[["delay", "year", "CI"]]
    f["Figure3a"] = (full_d[(full_d.reference_id == "Global:Demand-IP") & (full_d.delay <= 10)]
                     [["delay", "response", "penalty"]])
    f["Figure3b"] = cu[["delay", "year", "low_carbon_fuel_mass_equivalent_Mt_per_year"]]
    f["Figure4"] = pd.concat([
        g[(g.delay == 0) & (g.response == "Catch-up")],
        g[(g.delay == 5) & (g.response == "No catch-up")],
        g[(g.delay == 5) & (g.response == "Catch-up")],
    ])[["delay", "response", "year", "emissions", "cumulative_emissions"]]

    rows = []
    for family, refs in {"Demand": ["Demand-AR", "Demand-BAU", "Demand-IP"],
                         "Efficiency": ["EI-AR", "EI-BAU", "Demand-IP"]}.items():
        for ref in refs:
            for d in (3, 5, 10):
                r = full_d[(full_d.reference_id == f"Global:{ref}") & (full_d.delay == d) &
                           (full_d.response == "Catch-up")].iloc[0]
                rows.append(dict(family=family, display_case=ref.split("-")[1], delay=d, penalty=r.penalty))
    f["Figure5"] = pd.DataFrame(rows)

    r5 = g[g.delay == 5].copy()
    r5["overshoot_Gt"] = r5.cumulative_emissions - r5.cumulative_reference
    r5["overshoot_pct"] = 100 * r5.overshoot_Gt / r5.cumulative_reference
    f["Figure6_curves"] = pd.concat([r5[r5.response == "No catch-up"], r5[r5.response == "Catch-up"]])[
        ["response", "year", "overshoot_Gt", "overshoot_pct"]]
    f["Figure6_gates"] = gates

    f["FigureS1"] = (full_d[full_d.reference_id == "Australia:BITRE-AES-proxy"]
                     [["delay", "response", "penalty", "feasible"]])
    f["FigureS2"] = pd.concat([
        full_d[(full_d.reference_id == f"Australia:{c}") & full_d.delay.isin([3, 5, 10]) &
               (full_d.response == "Catch-up")].assign(case=c) for c in ("BAU", "IP", "AR")
    ])[["case", "delay", "penalty"]]

    shape_ids = {"Source": "Global:Demand-IP", "Front-loaded": "Shape:Front-loaded", "Linear": "Shape:Linear",
                 "Back-loaded": "Shape:Back-loaded", "Logistic": "Shape:Logistic"}
    a3, b3 = [], []
    for name, ref in shape_ids.items():
        q = full_a[(full_a.reference_id == ref) & (full_a.delay == 0) & (full_a.response == "Catch-up")].sort_values("year")
        a3.append(pd.DataFrame(dict(shape=name, year=q.year, CI=q.CI,
                                    normalized_remaining_CI=(q.CI - q.CI.iloc[-1]) / (q.CI.iloc[0] - q.CI.iloc[-1]))))
        p = full_d[(full_d.reference_id == ref) & full_d.delay.isin([3, 5, 10]) &
                   (full_d.response == "Catch-up")].sort_values("delay")
        b3.append(pd.DataFrame(dict(case=name, delay=p.delay, penalty=p.penalty)))
    f["FigureS3a"] = pd.concat(a3)
    f["FigureS3b"] = pd.concat(b3)
    return {k: v.reset_index(drop=True) for k, v in f.items()}


def write_workbook(frames: dict[str, pd.DataFrame], path: Path) -> None:
    wb = Workbook()
    idx = wb.active
    idx.title = "Index"
    idx.append(["Sheet", "Contents"])
    for c in idx[1]:
        c.font = Font(bold=True)
    for name, df in frames.items():
        idx.append([name, SHEETS[name]])
        ws = wb.create_sheet(name)
        ws.append([SHEETS[name]])
        ws["A1"].font = Font(bold=True)
        ws.append(list(df.columns))
        for c in ws[2]:
            c.font = Font(bold=True)
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v)
                       for v in row])
        ws.freeze_panes = "A3"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(10, min(28, max(len(str(c.value or "")) for c in col[1:]) + 2))
    idx.column_dimensions["A"].width = 16
    idx.column_dimensions["B"].width = 120
    wb.properties.title = "Figure source data"
    wb.properties.creator = ""
    wb.save(path)


def read_workbook(path: Path) -> dict[str, pd.DataFrame]:
    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    for name in SHEETS:
        rows = list(wb[name].iter_rows(values_only=True))
        out[name] = pd.DataFrame(rows[2:], columns=rows[1])
    return out


if (RUN / "annual_results.csv").exists():
    F = build_from_run()
    write_workbook(F, RUN / "figure_data.xlsx")
    OUT = RUN / "figures"
    print("Figure data built from output/ and written to output/figure_data.xlsx")
else:
    F = read_workbook(SHIPPED)
    OUT = ROOT / "figures"
    print("output/ not found: using results/figure_data.xlsx")
OUT.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------------------------
# Figures
# ---------------------------------------------------------------------------
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.spines.top": False,
                     "axes.spines.right": False, "pdf.fonttype": 42, "axes.titleweight": "bold"})
COL = {0: "#333333", 3: "#3878AA", 5: "#00897B", 10: "#C47B22"}
RESP = {"No catch-up": "#C8453C", "Catch-up": "#00897B"}


def save(fig: plt.Figure, stem: str) -> None:
    for ext in ("pdf", "png"):
        fig.savefig(OUT / f"{stem}.{ext}", dpi=300, facecolor="white")
    plt.close(fig)


# Figure 2: catch-up pathways.
f2 = F["Figure2"]
fig, ax = plt.subplots(figsize=(9, 5.5))
fig.subplots_adjust(left=.11, right=.97, bottom=.14, top=.94)
for d in (0, 3, 5, 10):
    q = f2[f2.delay == d].sort_values("year")
    ax.plot(q.year, q.CI, color=COL[d], lw=2)
    if d:
        first = q[q.year == 2026 + d].iloc[0]
        ax.scatter(first.year, first.CI, facecolor="white", edgecolor=COL[d], s=42, lw=1.5, zorder=5)
        pace = 100 * (TT / (TT - d) - 1)
        offsets = {3: (-30, 31), 5: (55, 34), 10: (16, 25)}
        ax.annotate(f"{d}-year delay\n{pace:.1f}% faster", (first.year, first.CI), xytext=offsets[d],
                    textcoords="offset points", color=COL[d], fontsize=9, ha="center",
                    arrowprops=dict(arrowstyle="-", color=COL[d], lw=.6))
pt = f2[(f2.delay == 0) & (f2.year == 2038)].iloc[0]
ax.annotate("Reference", (pt.year, pt.CI), xytext=(-50, -32), textcoords="offset points", fontsize=9,
            color=COL[0], arrowprops=dict(arrowstyle="-", color=COL[0], lw=.6))
ax.set(xlabel="Year", ylabel="Fuel carbon intensity (g CO₂/MJ)", xlim=(2024.5, 2051), ylim=(0, 92))
save(fig, "Figure2_catchup_pathways")

# Figure 3: cumulative penalties and supply-equivalent trajectories.
fig, axs = plt.subplots(1, 2, figsize=(11, 4.8))
fig.subplots_adjust(bottom=.22, wspace=.30, top=.86, left=.08, right=.97)
for response, color in RESP.items():
    q = F["Figure3a"][F["Figure3a"].response == response].sort_values("delay")
    axs[0].plot(q.delay, q.penalty, color=color, lw=2, label=response)
axs[0].set(title="(a) Cumulative penalty", xlabel="Delay (years)", ylabel="Cumulative penalty (Gt CO₂)")
axs[0].legend(frameon=False)
peaks = []
for d in (0, 3, 5, 10):
    q = F["Figure3b"][F["Figure3b"].delay == d].sort_values("year")
    m = q.low_carbon_fuel_mass_equivalent_Mt_per_year
    peaks.append(np.diff(m.to_numpy(float)).max())
    axs[1].plot(q.year, m, color=COL[d], lw=2, label="Reference" if d == 0 else f"{d}-year catch-up")
axs[1].set(title="(b) Annual low-carbon-fuel-equivalent supply", xlabel="Year",
           ylabel="Fuel mass equivalent (Mt/yr, at 44 MJ/kg)")
axs[1].legend(frameon=False, fontsize=8)
fig.text(.08, .045, "Peak year-on-year additions (reference / 3 / 5 / 10-year CU): " +
         " / ".join(f"{x:.1f}" for x in peaks) + " Mt/yr added over one-year steps.", fontsize=8.5)
save(fig, "Figure3_penalty_and_fuel_equivalent")

# Figure 4: annual and cumulative emissions after a five-year delay.
f4 = F["Figure4"]
fig, axs = plt.subplots(1, 2, figsize=(10.5, 4.8))
fig.subplots_adjust(bottom=.20, wspace=.32, top=.84, left=.08, right=.97)
for d, response, label, color in [(0, "Catch-up", "Reference", "#333333"),
                                  (5, "Catch-up", "Catch-up", RESP["Catch-up"]),
                                  (5, "No catch-up", "No catch-up", RESP["No catch-up"])]:
    q = f4[(f4.delay == d) & (f4.response == response)].sort_values("year")
    for ax, column in zip(axs, ("emissions", "cumulative_emissions")):
        ax.plot(q.year, q[column], label=label, color=color, lw=2)
for ax in axs:
    ax.axvline(2030, color="#999999", ls=":", lw=1)
    ax.set_xlabel("Year"); ax.set_ylim(bottom=0); ax.set_xlim(2025, 2050)
    ax.text(2030, .98, "End of delay", rotation=90, va="top", ha="right", fontsize=8, color="#666666",
            transform=ax.get_xaxis_transform())
axs[0].set(title="(a) Annual emissions", ylabel="Gt CO₂/yr")
axs[1].set(title="(b) Cumulative emissions", ylabel="Gt CO₂")
end = f4[f4.year == 2050].set_index(["delay", "response"]).cumulative_emissions
pen5 = end[(5, "Catch-up")] - end[(0, "Catch-up")]
axs[1].text(.34, .16, f"Catch-up penalty: {pen5:.1f} Gt CO₂\n({100 * pen5 / end[(0, 'Catch-up')]:.1f}% of reference total)",
            transform=axs[1].transAxes, fontsize=8.5, color=RESP["Catch-up"],
            bbox=dict(facecolor="white", edgecolor="none", alpha=.9))
fig.legend(*axs[1].get_legend_handles_labels(), loc="upper center", bbox_to_anchor=(.53, 1.0), ncol=3, frameon=False)
fig.text(.08, .045, "Five-year full pause through 2030; improvement resumes in 2031. "
         "Cumulative totals sum annual observations.", fontsize=8.5)
save(fig, "Figure4_global_emissions")

# Figure 5: demand and energy-intensity sensitivities.
f5 = F["Figure5"]
fig, axs = plt.subplots(1, 2, figsize=(10.5, 4.8))
fig.subplots_adjust(bottom=.18, wspace=.30, top=.86, left=.08, right=.97)
for ax, family in zip(axs, ("Demand", "Efficiency")):
    for j, d in enumerate((3, 5, 10)):
        vals = [f5[(f5.family == family) & (f5.display_case == c) & (f5.delay == d)].penalty.iloc[0]
                for c in ("BAU", "IP", "AR")]
        ax.bar(np.arange(3) + (j - 1) * .24, vals, .24, color=COL[d], label=f"{d}-year delay")
    ax.set_xticks(range(3), ["BAU", "IP", "AR"])
    ax.set(title=f"{family} sensitivity", ylabel="Catch-up penalty (Gt CO₂)")
    ax.legend(frameon=False, fontsize=8)
save(fig, "Figure5_global_sensitivity")

# Figure 6: policy-review diagnostics.
cur, gates = F["Figure6_curves"], F["Figure6_gates"]
fig = plt.figure(figsize=(12, 8.5))
ax = fig.add_axes([.105, .49, .77, .37])
fig.text(.045, .96, "(a) Cumulative emissions overshoot to date", fontsize=14, weight="bold")
crossing = {}
for response, color in RESP.items():
    q = cur[cur.response == response].sort_values("year")
    crossing[response] = int(q.year[q.overshoot_pct > 5].min())
    ax.plot(q.year, q.overshoot_pct, color=color, lw=2, ls="--" if response == "No catch-up" else "-",
            label="Catch-up from 2031" if response == "Catch-up" else "No catch-up")
    ax.text(2050.25, q.overshoot_pct.iloc[-1], f"{q.overshoot_pct.iloc[-1]:.1f}%", color=color, va="center", weight="bold")
ax.axhline(5, color="#888888", ls=":", label="Illustrative review signal: >5.0%")
for _, row in gates.iterrows():
    year = int(row.year)
    ax.axvline(year, color="#DDDDDD", ls=":", zorder=0)
    ax.text(year, 1.025, f"{year} review", ha="center", transform=ax.get_xaxis_transform(), fontsize=9)
    if year != 2030:
        ax.annotate(f"{row.overshoot_NC_pct:.1f}%", (year, row.overshoot_NC_pct), xytext=(5, 12),
                    textcoords="offset points", color=RESP["No catch-up"], fontsize=9)
    ax.annotate(f"{row.overshoot_CU_pct:.1f}%", (year, row.overshoot_CU_pct), xytext=(5, -18),
                textcoords="offset points", color=RESP["Catch-up"], fontsize=9)
ax.set(xlim=(2025, 2052.8), ylim=(0, cur.overshoot_pct.max() * 1.12), xlabel="Year",
       ylabel="Cumulative emissions overshoot\n(% of reference emissions to date)")
ax.legend(loc="lower center", bbox_to_anchor=(.5, 1.115), ncol=3, frameon=False, fontsize=9)
fig.text(.105, .425, f"First annual >5% crossing — No catch-up: {crossing['No catch-up']}; "
         f"Catch-up: {crossing['Catch-up']}", fontsize=9, color="#555555")
fig.text(.045, .375, "(b) Recovery and delivery pressure at review gates", fontsize=14, weight="bold")
fig.text(.055, .348, "Pace and supply: no catch-up before each review, then immediate endpoint-restoring recovery.", fontsize=10)
tax = fig.add_axes([.045, .125, .91, .205]); tax.axis("off")
headers = ["Review\ngate", "Overshoot:\ncatch-up (%)", "Overshoot:\nno catch-up (%)",
           "Required\ncatch-up pace\n(% faster)", "Peak supply-equivalent\naddition\n(Mt yr⁻¹)", "Signals\ncrossed"]
rows = [[str(int(r.year)), f"{r.overshoot_CU_pct:.1f}%\n({r.overshoot_CU_Gt:.1f} Gt)",
         f"{r.overshoot_NC_pct:.1f}%\n({r.overshoot_NC_Gt:.1f} Gt)", f"{r.required_pace_faster_pct:.1f}",
         f"{r.review_peak_Mt_per_year:.1f} vs {r.reference_peak_remaining_Mt_per_year:.1f}",
         str(r.trigger).replace(" + ", " +\n")] for _, r in gates.iterrows()]
table = tax.table(cellText=rows, colLabels=headers, cellLoc="center",
                  colWidths=[.07, .195, .195, .175, .195, .17], bbox=[0, 0, 1, 1])
table.auto_set_font_size(False); table.set_fontsize(9)
for (r, _c), cell in table.get_celld().items():
    cell.set_edgecolor("#CCCCCC"); cell.set_linewidth(.5)
    if r == 0:
        cell.set_facecolor("#EAF0F4"); cell.set_text_props(weight="bold")
    elif r % 2:
        cell.set_facecolor("#F5F7F9")
fig.text(.055, .082, "Parentheses: cumulative excess CO₂ (Gt). Supply entries: recovery vs reference over the same remaining-year window.", fontsize=8.3)
fig.text(.055, .055, "Supply additions are year-on-year fuel-mass-equivalent increases at 44 MJ kg⁻¹, not physical SAF tonnage.", fontsize=8.3)
fig.text(.055, .033, "Illustrative signals: overshoot >5.0%; required pace >25.0% faster. Equality does not cross the pace signal.", fontsize=8.3)
save(fig, "Figure6_policy_review_gates")

# Figure S1: Australian delay grid.
s1 = F["FigureS1"]
fig, ax = plt.subplots(figsize=(8, 4.8)); fig.subplots_adjust(bottom=.22)
for response, color in RESP.items():
    q = s1[(s1.response == response) & (s1.feasible.astype(bool))].sort_values("delay")
    ax.plot(q.delay, q.penalty, color=color, lw=2, label=response)
ax.set(xlabel="Delay (years)", ylabel="Cumulative penalty (Mt CO₂)",
       title="Australian domestic aviation: explicit 2025 energy proxy")
ax.legend(frameon=False)
fig.text(.12, .05, "AES FY2023–24: 135.251 PJ assigned to 2025; BITRE projected activity. "
         "Sector boundaries are not fully matched.", fontsize=8.5)
save(fig, "FigureS1_australia_delay_curve")

# Figure S2: Australian demand sensitivity.
s2 = F["FigureS2"]
fig, ax = plt.subplots(figsize=(8, 4.8)); fig.subplots_adjust(bottom=.22)
for j, d in enumerate((3, 5, 10)):
    vals = [s2[(s2.case == c) & (s2.delay == d)].penalty.iloc[0] for c in ("BAU", "IP", "AR")]
    ax.bar(np.arange(3) + (j - 1) * .24, vals, .24, color=COL[d], label=f"{d}-year delay")
ax.set_xticks(range(3), ["BAU", "IP", "AR"])
ax.set(ylabel="Catch-up penalty (Mt CO₂)", title="Australian demand sensitivity")
ax.legend(frameon=False)
fig.text(.12, .05, "Common AES/BITRE 2025-proxy efficiency and 73.5 g/MJ starting intensity; harmonised demand variants.", fontsize=8.5)
save(fig, "FigureS2_australia_demand_sensitivity")

# Figure S3: alternative transition shapes.
fig, axs = plt.subplots(1, 2, figsize=(11, 4.8)); fig.subplots_adjust(bottom=.17, wspace=.28, top=.86)
for name, color in zip(["Source", "Front-loaded", "Linear", "Back-loaded", "Logistic"],
                       ["#333333", "#3878AA", "#00897B", "#C47B22", "#8865AD"]):
    q = F["FigureS3a"][F["FigureS3a"]["shape"] == name].sort_values("year")
    axs[0].plot(q.year, q.normalized_remaining_CI, label=name, color=color, lw=2)
    p = F["FigureS3b"][F["FigureS3b"]["case"] == name].sort_values("delay")
    axs[1].plot(p.delay, p.penalty, color=color, marker="o", lw=2, label=name)
axs[0].set(title="(a) Reference transition shape", xlabel="Year", ylabel="Remaining normalized carbon intensity")
axs[0].legend(frameon=False, fontsize=8)
axs[1].set(title="(b) Catch-up penalties", xlabel="Delay (years)", ylabel="Gt CO₂", xticks=[3, 5, 10])
save(fig, "FigureS3_transition_shape")

print(f"Figures written to {OUT}")
